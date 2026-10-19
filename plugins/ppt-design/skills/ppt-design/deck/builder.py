# -*- coding: utf-8 -*-
"""
Deck — 덱 빌더. 테마(폰트+스타일)와 덱 유형을 받아 슬라이드를 만든다.
16:9 (1280x720 px). 출력은 <PPT_DESIGN_OUTPUT 또는 cwd/OUTPUT>/ppt-design/<프로젝트>/.
"""
import os
import tempfile
from pptx import Presentation
import tokens as T
import theme as theme_mod
from helpers import px, rgb, text


def output_root():
    # PPT_DESIGN_OUTPUT 이 있으면 그걸, 없으면 현재 작업폴더/OUTPUT
    return os.environ.get("PPT_DESIGN_OUTPUT") or os.path.join(os.getcwd(), "OUTPUT")


def project_dir(project, sub=None):
    # 결과물은 <root>/ppt-design/<project>[/<sub>] — 스킬별로 묶는다
    parts = [output_root(), "ppt-design", project] + ([sub] if sub else [])
    d = os.path.join(*parts)
    os.makedirs(d, exist_ok=True)
    return d


class Deck:
    def __init__(self, theme=None, kind=None, brand=""):
        """theme: 프리셋 이름/파일경로/dict (생략 시 기본 cream-coral).
        kind: 'proposal'|'report'|'lecture'|'guidebook' (푸터·계획 가이드용)."""
        if theme is not None:
            theme_mod.use(theme)
        self.kind = kind
        self.brand = brand
        # 사진 레이아웃이 자르거나 둥글게 깎은 그림을 두는 곳 (작업 폴더를 어지럽히지 않게)
        self.img_cache = os.path.join(tempfile.gettempdir(), "ppt-design-imgcache")
        self.prs = Presentation()
        self.prs.slide_width = px(T.SLIDE_W_PX)
        self.prs.slide_height = px(T.SLIDE_H_PX)
        self.blank = self.prs.slide_layouts[6]
        self._page = 0

    def add(self, bg="canvas"):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        return s

    def footer(self, slide, dark=False, label=None, color=None):
        """바닥글(브랜드 + 쪽 번호). color 를 주면 그 색으로(색 배경 슬라이드용)."""
        self._page += 1
        c = color or ("on_dark_soft" if dark else "muted_soft")
        lab = self.brand if label is None else label
        if lab:
            text(slide, 80, 682, 500, 26, lab, type_token="body_sm",
                 color=c, valign="middle", size_px=11, line_spacing=1.2)
        text(slide, T.SLIDE_W_PX - 80 - 60, 682, 60, 26, f"{self._page:02d}",
             type_token="body_sm", color=c, align="right", valign="middle",
             size_px=11, line_spacing=1.2)

    def _korean_word_wrap(self):
        """한글이 글자 단위로 끊기지 않게(«사진이에» / «요.») 한글이 든 글자에
        한국어 표시(lang="ko-KR")를 붙인다. 파워포인트는 이 표시가 있어야 한글을
        단어 단위로 줄바꿈한다. 문단 설정(eaLnBrk·latinLnBrk)으로는 바뀌지 않았다."""
        import re
        hangul = re.compile("[가-힣ㄱ-ㆎ]")
        def shapes(container):
            for sh in container:
                if sh.shape_type == 6:          # 그룹
                    yield from shapes(sh.shapes)
                else:
                    yield sh

        def frames(sh):
            if sh.has_text_frame:
                yield sh.text_frame
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

        for slide in self.prs.slides:
            for sh in shapes(slide.shapes):
                for tf in frames(sh):
                    for p in tf.paragraphs:
                        for r in p.runs:
                            if hangul.search(r.text):
                                r._r.get_or_add_rPr().set("lang", "ko-KR")

    def save(self, path):
        self._korean_word_wrap()
        self.prs.save(path)
        return path

    def save_project(self, project, filename="deck.pptx"):
        """<root>/ppt-design/<project>/<filename> 에 저장 (root=PPT_DESIGN_OUTPUT 또는 cwd/OUTPUT)."""
        path = os.path.join(project_dir(project), filename)
        self._korean_word_wrap()
        self.prs.save(path)
        return path
